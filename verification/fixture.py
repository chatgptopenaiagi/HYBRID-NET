"""Build the frozen Genesis 2 DOCX. This generates input, never expected output."""
from __future__ import annotations

import argparse
from datetime import datetime, timezone
import hashlib
import io
import json
from pathlib import Path
import zipfile


def create_fixture(output: Path) -> dict:
    from docx import Document
    from docx.shared import Mm, Pt
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    output = output.resolve()
    criteria_source = Path(__file__).with_name("genesis2-criteria.json")
    criteria_bytes = criteria_source.read_bytes()
    spec = json.loads(criteria_bytes)
    criteria_output = output.with_name("criteria.json")
    if output.suffix.lower() != ".docx":
        raise ValueError("Output must name a .docx file")
    if output.exists() or criteria_output.exists():
        raise FileExistsError("Fixture and sibling criteria.json must be new paths")
    output.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Mm(210), Mm(297)
    section.top_margin = section.bottom_margin = Mm(20)
    section.left_margin = section.right_margin = Mm(20)
    style = doc.styles["Normal"]
    style.font.name = "Arial"
    style.font.size = Pt(11)
    style.paragraph_format.line_spacing = 1.15
    style.paragraph_format.space_after = Pt(8)
    style.paragraph_format.widow_control = False
    style.element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
    properties = doc.core_properties
    properties.title = spec["title"]
    properties.author = "HYBRID.NET controlled fixture"
    properties.last_modified_by = "HYBRID.NET controlled fixture"
    properties.created = properties.modified = datetime(2026, 10, 22, tzinfo=timezone.utc)
    properties.revision = 1
    title = doc.add_paragraph()
    title_run = title.add_run(spec["title"])
    title_run.bold, title_run.font.size = True, Pt(16)
    for text in spec["paragraphs"]["1"]:
        doc.add_paragraph(text)
    table = doc.add_table(rows=4, cols=3)
    table.style = "Table Grid"
    table.autofit = False
    for column, width in zip(table.columns, spec["layout"]["table_columns_mm"]):
        column.width = Mm(width)
    for row, values in zip(table.rows, spec["table_matrix"]):
        for cell, text, width in zip(row.cells, values, spec["layout"]["table_columns_mm"]):
            cell.width, cell.text = Mm(width), text
        cant_split = OxmlElement("w:cantSplit")
        row._tr.get_or_add_trPr().append(cant_split)
    doc.add_page_break()
    heading = doc.add_paragraph()
    heading.add_run(spec["second_heading"]).bold = True
    doc.add_paragraph(spec["paragraphs"]["2"][0])
    doc.add_paragraph(spec["closing_line"])
    memory = io.BytesIO()
    doc.save(memory)
    # Stable ZIP timestamps/order make repeated generation byte-identical.
    target = io.BytesIO()
    with zipfile.ZipFile(memory) as original, zipfile.ZipFile(target, "w", zipfile.ZIP_DEFLATED) as archive:
        for name in sorted(original.namelist()):
            entry = zipfile.ZipInfo(name, (2026, 10, 22, 0, 0, 0))
            entry.compress_type = zipfile.ZIP_DEFLATED
            entry.external_attr = 0o600 << 16
            archive.writestr(entry, original.read(name))
    data = target.getvalue()
    with output.open("xb") as stream:
        stream.write(data)
    with criteria_output.open("xb") as stream:
        stream.write(criteria_bytes)
    return {"fixture_id": spec["criteria_id"], "fixture_version": "1.0", "source": str(output),
            "source_sha256": hashlib.sha256(data).hexdigest(), "size": len(data),
            "criteria_path": str(criteria_output), "criteria_sha256": hashlib.sha256(criteria_bytes).hexdigest(),
            "expected_page_count": 2, "layout": spec["layout"]}


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--output", type=Path, required=True)
    args = parser.parse_args()
    try:
        result = create_fixture(args.output)
    except (OSError, ValueError) as exc:
        print(json.dumps({"error": type(exc).__name__, "detail": str(exc)}))
        return 2
    print(json.dumps(result, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
